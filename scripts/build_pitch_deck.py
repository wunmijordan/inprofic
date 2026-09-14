#!/usr/bin/env python3
"""Build the editable INPROFIC product pitch deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "docs" / "INPROFIC_Pitch_Deck.pptx"
BRAND_DIR = ROOT / "apps" / "core" / "static" / "core" / "brand"
LOGO_ON_DARK = BRAND_DIR / "inprofic-logo-on-dark.png"
LOGO_ON_LIGHT = BRAND_DIR / "inprofic-logo-on-light.png"

NAVY = "050733"
ORANGE = "D14900"
PAPER = "FAF6EF"
WHITE = "FFFFFF"
INK = "211D1A"
MUTED = "6B6662"
GOLD = "E7BD65"
GREEN = "18794E"
PALE = "F5EFE0"
LINE = "DED5C4"


def rgb(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, color, radius=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.color.rgb = rgb(line or color)
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False,
         font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.01)
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    paragraph.alignment = align
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = rgb(color)
    return box


def rich_text(slide, runs, x, y, w, h, size=18, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    for value, run_color, bold in runs:
        run = paragraph.add_run()
        run.text = value
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(run_color or color)
    return box


def base_slide(prs, title_value, kicker=None, dark=False, number=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(NAVY if dark else PAPER)
    if kicker:
        text(slide, kicker.upper(), 0.72, 0.38, 8.8, 0.25, 10,
             GOLD if dark else ORANGE, True)
    text(slide, title_value, 0.7, 0.72, 11.9, 0.72, 29,
         WHITE if dark else INK, True)
    if number is not None:
        text(slide, f"{number:02d}", 12.15, 7.05, 0.48, 0.2, 9,
             "A8A7C3" if dark else MUTED, False, align=PP_ALIGN.RIGHT)
    slide.shapes.add_picture(
        str(LOGO_ON_DARK if dark else LOGO_ON_LIGHT),
        Inches(0.72), Inches(6.98), width=Inches(1.22),
    )
    return slide


def card(slide, title_value, body, x, y, w, h, accent=ORANGE, dark=False):
    rect(slide, x, y, w, h, "111347" if dark else WHITE, True,
         "292B62" if dark else LINE)
    rect(slide, x, y, 0.08, h, accent)
    text(slide, title_value, x + 0.28, y + 0.24, w - 0.5, 0.36, 16,
         WHITE if dark else INK, True)
    text(slide, body, x + 0.28, y + 0.72, w - 0.5, h - 0.9, 11.5,
         "C7C8DD" if dark else MUTED)


def pill(slide, value, x, y, w, dark=False):
    rect(slide, x, y, w, 0.38, "181A55" if dark else PALE, True)
    text(slide, value, x + 0.08, y + 0.09, w - 0.16, 0.16, 9.5,
         WHITE if dark else INK, True, align=PP_ALIGN.CENTER)


def arrow(slide, x, y, w, color=ORANGE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.35))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.color.rgb = rgb(color)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 — Cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(NAVY)
    rect(slide, 8.75, -0.4, 5.2, 8.2, "0B0D45", True)
    slide.shapes.add_picture(str(LOGO_ON_DARK), Inches(0.8), Inches(0.65), width=Inches(3.15))
    rect(slide, 0.8, 1.55, 1.2, 0.08, ORANGE)
    text(slide, "From inventory to commerce—\none connected operation.", 0.8, 2.0, 7.2, 1.65, 34, WHITE, True)
    text(slide, "A production-aware commercial management system for businesses that buy, make, stock, sell and account for value.", 0.82, 4.1, 6.9, 1.0, 16, "C7C8DD")
    text(slide, "PRODUCT PITCH DECK", 0.82, 6.65, 3.0, 0.25, 10, GOLD, True)
    slide.shapes.add_picture(
        str(BRAND_DIR / "inprofic-mark-on-dark.png"),
        Inches(9.55), Inches(1.45), width=Inches(3.75),
    )

    # 2 — Problem
    slide = base_slide(prs, "Growing operations fracture across disconnected tools.", "The problem", number=2)
    card(slide, "Stock lives alone", "Purchases, usage, finished stock and availability drift when every team maintains a different sheet.", 0.72, 1.72, 3.75, 3.0)
    card(slide, "Production loses context", "Customer demand, recipes, planned output, wastage and shortages become difficult to reconcile.", 4.78, 1.72, 3.75, 3.0, GOLD)
    card(slide, "Money arrives late", "Sales, receivables, payables, expenses and cash balances are reconstructed after the operational event.", 8.84, 1.72, 3.75, 3.0, GREEN)
    text(slide, "The cost is not another spreadsheet. It is delayed decisions, weak traceability and preventable leakage.", 1.3, 5.35, 10.7, 0.7, 20, NAVY, True, align=PP_ALIGN.CENTER)

    # 3 — Solution chain
    slide = base_slide(prs, "One chain of truth from purchase to payment.", "The INPROFIC answer", dark=True, number=3)
    steps = [("PROCURE", "Supplier orders\nand arrivals"), ("STOCK", "Materials and\nfinished goods"), ("PRODUCE", "Plans, batches,\nQC and yield"), ("SELL", "Counter, online\nand distribution"), ("ACCOUNT", "Cash, credit,\nexpenses and margin")]
    x = 0.7
    for index, (heading, body) in enumerate(steps):
        rect(slide, x, 2.05, 2.15, 2.45, "111347", True, "292B62")
        text(slide, f"0{index + 1}", x + 0.2, 2.28, 0.4, 0.25, 11, GOLD, True)
        text(slide, heading, x + 0.2, 2.84, 1.75, 0.35, 15, WHITE, True)
        text(slide, body, x + 0.2, 3.38, 1.75, 0.65, 11, "C7C8DD")
        if index < len(steps) - 1:
            arrow(slide, x + 2.18, 3.06, 0.42)
        x += 2.5
    text(slide, "Each event updates the next layer without losing the business, customer, product or audit context.", 1.25, 5.4, 10.8, 0.55, 18, WHITE, True, align=PP_ALIGN.CENTER)

    # 4 — Acronym
    slide = base_slide(prs, "The name carries the operating model.", "Why INPROFIC", number=4)
    parts = [("IN", "Inventory", "Know what exists and how it moves."), ("PRO", "Procurement + Production", "Connect inputs, plans and physical output."), ("FI", "Finance", "Follow the money created by operations."), ("C", "Commerce", "Turn validated demand into controlled fulfilment.")]
    positions = [(0.72, 1.65), (6.78, 1.65), (0.72, 4.08), (6.78, 4.08)]
    for (prefix, heading, body), (x, y) in zip(parts, positions):
        rect(slide, x, y, 5.82, 1.85, WHITE, True, LINE)
        text(slide, prefix, x + 0.28, y + 0.34, 1.05, 0.58, 25, ORANGE, True)
        text(slide, heading, x + 1.42, y + 0.3, 4.0, 0.38, 16, NAVY, True)
        text(slide, body, x + 1.42, y + 0.84, 3.95, 0.55, 11.5, MUTED)

    # 5 — Verticals
    slide = base_slide(prs, "One core, vocabulary shaped to the business.", "Built for five operating realities", dark=True, number=5)
    verticals = [("BAKERY", "Recipes, batches, offcuts and distribution"), ("RESTAURANT", "Recipes, preparation, table service and bulk demand"), ("PRODUCTION", "Formulas, bills of materials and controlled output"), ("WHOLESALE", "Supplier arrivals, distribution pricing and customer credit"), ("RETAIL", "Stock-first purchasing and multichannel sales")]
    for i, (heading, body) in enumerate(verticals[:3]):
        x = 0.72 + i * 4.1
        card(slide, heading, body, x, 1.65, 3.75, 1.62, GOLD if i == 1 else ORANGE, True)
    card(slide, verticals[3][0], verticals[3][1], 0.72, 3.8, 5.78, 1.62, GREEN, True)
    card(slide, verticals[4][0], verticals[4][1], 6.78, 3.8, 5.78, 1.62, ORANGE, True)

    # 6 — Inventory/procurement
    slide = base_slide(prs, "Inventory and procurement speak the same unit language.", "Control inputs", number=6)
    card(slide, "Three-layer units", "Buy by bag or carton, track the package in kg or litres, and consume recipes in grams, spoons or caps—without guessing conversions.", 0.72, 1.65, 3.75, 3.45)
    card(slide, "Arrival-driven stock", "Receiving a purchase updates stock, landed unit cost and supplier history together.", 4.78, 1.65, 3.75, 3.45, GOLD)
    card(slide, "Movement visibility", "Low-stock warnings, movement history, market stock and finished inventory remain traceable by business.", 8.84, 1.65, 3.75, 3.45, GREEN)
    pill(slide, "RAW MATERIALS", 1.5, 5.55, 2.15); arrow(slide, 3.82, 5.56, 0.72)
    pill(slide, "PURCHASE ARRIVAL", 4.72, 5.55, 2.2); arrow(slide, 7.08, 5.56, 0.72)
    pill(slide, "AVAILABLE STOCK", 8.0, 5.55, 2.2); arrow(slide, 10.36, 5.56, 0.72)
    pill(slide, "COST HISTORY", 11.18, 5.55, 1.45)

    # 7 — Production
    slide = base_slide(prs, "Production keeps demand, plans and physical truth distinct.", "Production-aware by design", dark=True, number=7)
    items = [("Proportional recipes", "Material usage scales against registered batch yield using Decimal precision."), ("Shared Production Runs", "Coordinate multiple normal orders, combine stock checks and retain order-level costing."), ("Completion truth", "Record gross output, wastage, saleable units, shortage, planned offcut and unexpected excess."), ("Safe reconciliation", "Use traceable surplus or live stock without creating a duplicate sale or production event."), ("Compensating reversal", "Reverse auditable effects instead of deleting history, then edit as a new pending order.")]
    for i, (heading, body) in enumerate(items):
        y = 1.55 + i * 0.98
        text(slide, f"0{i+1}", 0.8, y + 0.02, 0.4, 0.25, 11, GOLD, True)
        text(slide, heading, 1.35, y, 3.0, 0.3, 15, WHITE, True)
        text(slide, body, 4.42, y, 7.65, 0.55, 11.5, "C7C8DD")
        rect(slide, 1.35, y + 0.68, 10.72, 0.015, "292B62")

    # 8 — Sales/commerce
    slide = base_slide(prs, "Commerce validates demand before it touches operations.", "Sell across channels", number=8)
    card(slide, "Physical Store", "Immediate stock-aware selling for walk-ins and direct transactions.", 0.72, 1.62, 3.75, 1.72)
    card(slide, "Online", "Published catalogues, baskets, channel pricing, payment and saved order tracking.", 4.78, 1.62, 3.75, 1.72, GOLD)
    card(slide, "Distribution", "Bulk minimums, customer/channel price rules, credit and planned production routes.", 8.84, 1.62, 3.75, 1.72, GREEN)
    text(slide, "PAYMENT-FIRST CHECKOUT", 0.74, 4.0, 3.3, 0.3, 12, ORANGE, True)
    stages = ["Validate basket", "Snapshot price", "Verify payment", "Create intake once", "Accept into operations"]
    for i, stage in enumerate(stages):
        x = 0.74 + i * 2.45
        pill(slide, stage, x, 4.62, 2.02)
        if i < 4: arrow(slide, x + 2.05, 4.63, 0.34)
    text(slide, "Hosted storefront • Order Now link • Headless API • Signed platform connector", 1.2, 5.65, 10.9, 0.5, 17, NAVY, True, align=PP_ALIGN.CENTER)

    # 9 — Finance
    slide = base_slide(prs, "The financial record follows the operational event.", "Finance without reconstruction", dark=True, number=9)
    card(slide, "Cash accounts", "Maintain actual money balances across cash, bank and card accounts.", 0.72, 1.65, 3.75, 2.05, ORANGE, True)
    card(slide, "Receivables & payables", "Keep customer credit and supplier obligations distinct from fulfilment status.", 4.78, 1.65, 3.75, 2.05, GOLD, True)
    card(slide, "Expenses", "Capture operational spend and the corresponding payment record.", 8.84, 1.65, 3.75, 2.05, GREEN, True)
    text(slide, "Sales", 1.0, 4.6, 1.1, 0.3, 16, WHITE, True); arrow(slide, 2.0, 4.6, 0.8)
    text(slide, "Financial transaction", 3.0, 4.6, 2.4, 0.3, 16, WHITE, True); arrow(slide, 5.42, 4.6, 0.8)
    text(slide, "Cash / credit state", 6.42, 4.6, 2.2, 0.3, 16, WHITE, True); arrow(slide, 8.7, 4.6, 0.8)
    text(slide, "Reports & margin", 9.72, 4.6, 2.1, 0.3, 16, WHITE, True)
    text(slide, "Corrections remain compensating and auditable—not silent rewrites.", 1.5, 5.7, 10.3, 0.5, 19, GOLD, True, align=PP_ALIGN.CENTER)

    # 10 — Controls
    slide = base_slide(prs, "Controls are part of the workflow, not an afterthought.", "Trust and accountability", number=10)
    card(slide, "Tenant isolation", "Every business-owned record stays scoped to the active workspace; posted business IDs are never trusted.", 0.72, 1.65, 3.75, 3.35)
    card(slide, "Role-aware access", "Business Admins, configurable roles and module permissions govern visibility; the Founder-only Demo role can explore live workflows without writes.", 4.78, 1.65, 3.75, 3.35, GOLD)
    card(slide, "Audit-preserving history", "Creator trails, movement records, immutable price/usage snapshots and reversal states retain what actually happened.", 8.84, 1.65, 3.75, 3.35, GREEN)
    text(slide, "Built to answer: who did what, for which business, to which stock and money record—and how was it corrected?", 1.1, 5.55, 11.1, 0.65, 18, NAVY, True, align=PP_ALIGN.CENTER)

    # 11 — Reporting
    slide = base_slide(prs, "Operational reporting stays close to the source data.", "Decision support", dark=True, number=11)
    reports = ["Inventory health", "Procurement spend", "Production performance", "Sales & channels", "Finance movements", "CSV / XLSX / JSON backup"]
    for i, report in enumerate(reports):
        x = 0.78 + (i % 3) * 4.12
        y = 1.62 + (i // 3) * 2.15
        rect(slide, x, y, 3.72, 1.7, "111347", True, "292B62")
        text(slide, f"↗  {report}", x + 0.25, y + 0.38, 3.2, 0.35, 16, WHITE, True)
        text(slide, "Same business scope and status logic as the operating screen.", x + 0.25, y + 0.94, 3.15, 0.42, 10.5, "C7C8DD")

    # 12 — Platform architecture
    slide = base_slide(prs, "Practical architecture, deployable without lock-in.", "Platform", number=12)
    layers = [("EXPERIENCE", "Django server-rendered workspace • public storefront • API and connector boundaries"), ("OPERATING CORE", "Accounts • Inventory • Procurement • Production • Sales • Expenses • Commerce • Reports"), ("DATA & FILES", "PostgreSQL/Supabase or SQLite • Cloudflare R2 or local media • WhiteNoise static assets"), ("RUNTIME", "PythonAnywhere WSGI or Render ASGI/Daphne • WebSocket notifications • health and scheduled jobs")]
    for i, (heading, body) in enumerate(layers):
        y = 1.55 + i * 1.22
        rect(slide, 1.1, y, 11.1, 0.93, WHITE, True, LINE)
        text(slide, heading, 1.38, y + 0.22, 1.65, 0.25, 11, ORANGE if i < 2 else GREEN, True)
        text(slide, body, 3.12, y + 0.18, 8.65, 0.45, 11.5, MUTED)
    text(slide, "One codebase. Environment-selected infrastructure. Portable operating data.", 1.5, 6.45, 10.3, 0.4, 17, NAVY, True, align=PP_ALIGN.CENTER)

    # 13 — Commercial model
    slide = base_slide(prs, "A modular path from first control to full operations.", "Subscription model", dark=True, number=13)
    plans = [("STARTER", "Core workspace and Basic Reports", "Begin with visibility"), ("PRODUCTION", "Production capabilities and Full Reports", "Control making and yield"), ("BUSINESS PRO", "All modules, including Finance and Commerce", "Run the connected business")]
    for i, (name, includes, promise) in enumerate(plans):
        x = 0.72 + i * 4.06
        rect(slide, x, 1.65, 3.72, 3.8, "111347", True, ORANGE if i == 2 else "292B62")
        text(slide, name, x + 0.3, 2.0, 3.1, 0.4, 18, WHITE, True)
        rect(slide, x + 0.3, 2.63, 0.75, 0.06, ORANGE if i == 2 else GOLD)
        text(slide, includes, x + 0.3, 3.0, 3.05, 0.95, 13, "C7C8DD")
        text(slide, promise, x + 0.3, 4.55, 3.0, 0.35, 12, GOLD, True)
    text(slide, "30-day trial • founder-controlled pricing & timed promos • animated public campaign creative • monthly/yearly billing", 1.1, 6.15, 11.2, 0.42, 15, WHITE, True, align=PP_ALIGN.CENTER)

    # 14 — Value
    slide = base_slide(prs, "Why INPROFIC wins the operating conversation.", "Value proposition", number=14)
    values = [("CONNECTED", "Operational and financial truth in the same workflow."), ("SPECIFIC", "Production nuance without forcing every business into generic ERP language."), ("AUDITABLE", "Stock, payments and corrections retain accountable history."), ("ADAPTABLE", "Five verticals, modular entitlements, Founder-managed promo campaigns and two supported deployment paths."), ("COMMERCE-READY", "Validated, payment-first demand flows safely into operations."), ("PRACTICAL", "Server-rendered, approachable and designed for day-to-day teams.")]
    for i, (heading, body) in enumerate(values):
        x = 0.72 + (i % 3) * 4.06
        y = 1.55 + (i // 3) * 2.15
        card(slide, heading, body, x, y, 3.72, 1.65, ORANGE if i in (0, 4) else GOLD)

    # 15 — Close
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(NAVY)
    slide.shapes.add_picture(str(LOGO_ON_DARK), Inches(0.8), Inches(0.68), width=Inches(2.55))
    text(slide, "Know the stock.\nControl the work.\nFollow the money.\nServe the customer.", 0.8, 1.65, 7.5, 3.25, 38, WHITE, True)
    rect(slide, 0.82, 5.25, 2.1, 0.08, ORANGE)
    text(slide, "Inventory • Procurement • Production • Finance • Commerce", 0.82, 5.72, 7.2, 0.42, 15, GOLD, True)
    rect(slide, 9.25, 1.35, 3.05, 4.75, "111347", True, "292B62")
    text(slide, "THE NEXT STEP", 9.65, 1.85, 2.25, 0.25, 11, GOLD, True, align=PP_ALIGN.CENTER)
    text(slide, "Put one real business workflow through INPROFIC—from purchase or order to stock, fulfilment and finance.", 9.65, 2.65, 2.25, 1.75, 18, WHITE, True, align=PP_ALIGN.CENTER)
    text(slide, "See the connected record, not just the feature list.", 9.65, 4.75, 2.25, 0.65, 12, "C7C8DD", align=PP_ALIGN.CENTER)

    prs.core_properties.title = "INPROFIC Product Pitch Deck"
    prs.core_properties.subject = "Inventory, procurement, production, finance and commerce"
    prs.core_properties.author = "INPROFIC"
    prs.core_properties.keywords = "INPROFIC, inventory, procurement, production, finance, commerce"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
