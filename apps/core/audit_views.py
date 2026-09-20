from decimal import Decimal
from io import BytesIO
import re

from django.contrib.auth.decorators import login_required
from django.db.models import Sum
from django.http import HttpResponse
from django.contrib import messages
from django.core.exceptions import ValidationError
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_date

from commerce.models import DeliveryAssignment
from expenses.models import Expense
from inventory.models import FinishedGood, RawMaterial, RawMaterialMeasurementChange
from procurement.models import PurchaseOrder, RawMaterialCostSnapshot
from production.models import ProductionBatch, ProductionCostSnapshot, ProductionQualityCheck
from sales.models import CustomerPayment, Sale
from accounts.services import is_business_admin, user_has_permission
from accounts.platform_integrations import redact_disabled_integrations

from .models import AuditLog, AuditQuery, CashAccount, FinancialTransaction
from .services import audit


_AUDIT_DETAIL_LABELS = {
    "previous_status": "Previous status",
    "status": "New status",
    "driver_id": "Rider record",
    "delivery_id": "Delivery reference",
    "external_reference": "Provider reference",
    "proof_reference": "Proof of delivery",
    "resolution_note": "Resolution",
    "customer_name": "Customer",
    "provider": "Delivery provider",
    "provider_account": "Provider account",
    "tracking_number": "Tracking number",
    "category": "Category",
    "reason": "Reason",
    "source": "Source",
    "module": "Module",
    "severity": "Priority",
    "amount": "Amount",
    "quantity": "Quantity",
    "configured": "Configuration ready",
    "created": "Created items",
    "existing": "Already active",
    "error": "Outcome note",
}
_AUDIT_PRIVATE_MARKERS = {
    "secret", "token", "password", "credential", "signature", "header",
    "payload", "response", "api_key", "webhook_secret",
}

# The external auditor's event stream is deliberately narrower than the
# internal AuditLog table. Configuration, access-management and maintenance
# events remain recorded internally, but this workspace presents commercial
# and stock evidence only.
_EXTERNAL_BUSINESS_ACTIVITY_MODELS = frozenset({
    # Stock and product catalogue
    "RawMaterial", "FinishedGood", "ProductCategory",
    "OperationalSupplyDispense", "StockAdjustment", "MarketStockLot",
    "DistributionReturn",
    # Procurement and finance
    "PurchaseOrder", "SupplierPayment", "Sale", "CustomerPayment",
    "CashAccount", "FinancialTransaction", "Expense", "ExpensePayment",
    # Commerce, checkout, payment and delivery execution
    "StorefrontProduct", "CommerceCheckoutSession", "CommerceIntake",
    "CommercePayment", "CommercePaymentClaim", "CommercePaymentReceipt",
    "DeliveryAssignment", "DeliveryIssue",
})


def _audit_detail_value(value):
    if value is True:
        return "Yes"
    if value is False:
        return "No"
    if value is None or value == "":
        return "Not recorded"
    if isinstance(value, dict):
        return "Recorded securely"
    if isinstance(value, (list, tuple)):
        value = ", ".join(str(item) for item in value)
    text = str(redact_disabled_integrations(value))
    return f"{text[:177]}…" if len(text) > 180 else text


def _public_audit_details(metadata):
    if not isinstance(metadata, dict):
        return []
    rows = []
    for key, value in metadata.items():
        normalized = str(key).lower()
        if any(marker in normalized for marker in _AUDIT_PRIVATE_MARKERS):
            continue
        label = _AUDIT_DETAIL_LABELS.get(normalized, normalized.replace("_", " ").title())
        rows.append({"label": label, "value": _audit_detail_value(value)})
    return rows[:8]


def _prepare_audit_log(log):
    log.description = redact_disabled_integrations(log.description)
    log.action_label = (log.action or "Activity").replace("_", " ").title()
    model = (log.model_name or "Record").split(".")[-1]
    log.record_type_label = re.sub(r"(?<!^)(?=[A-Z])", " ", model)
    log.public_details = _public_audit_details(log.metadata)
    return log


def _period(request):
    today = timezone.localdate()
    date_to = parse_date(request.GET.get("date_to", "")) or today
    date_from = parse_date(request.GET.get("date_from", "")) or (date_to - timezone.timedelta(days=90))
    if date_from > date_to:
        date_from, date_to = date_to, date_from
    return date_from, date_to


def _audit_dataset(request, *, for_export=False):
    business = request.business
    date_from, date_to = _period(request)

    materials = list(RawMaterial.objects.filter(business=business).order_by("name"))
    products = list(
        FinishedGood.objects.filter(business=business)
        .select_related("product_category")
        .prefetch_related("recipe_items__raw_material", "production_materials__raw_material")
        .order_by("name")
    )
    for material in materials:
        material.audit_value = Decimal(material.stock or 0) * Decimal(material.cost_per_unit or 0)
    material_value = sum((m.audit_value for m in materials), Decimal("0"))
    product_value = sum((Decimal(g.stock or 0) * Decimal(g.est_cost or 0) for g in products), Decimal("0"))

    purchases = list(
        PurchaseOrder.objects.filter(business=business, date__range=(date_from, date_to))
        .prefetch_related("items__raw_material", "items__finished_good", "payments")
        .order_by("-date", "-id")
    )
    purchase_total = sum((po.total for po in purchases), Decimal("0"))
    supplier_paid = sum((sum((p.amount for p in po.payments.all()), Decimal("0")) for po in purchases), Decimal("0"))
    supplier_payable = max(Decimal("0"), purchase_total - supplier_paid)

    sales = list(
        Sale.objects.filter(business=business, date__range=(date_from, date_to))
        .prefetch_related("items__finished_good", "payments")
        .order_by("-date", "-id")
    )
    sales_total = sum((sale.total for sale in sales), Decimal("0"))
    sales_paid = sum((sum((p.amount for p in sale.payments.all()), Decimal("0")) for sale in sales), Decimal("0"))
    receivables = max(Decimal("0"), sales_total - sales_paid)

    expenses = list(Expense.objects.filter(business=business, date__range=(date_from, date_to)).order_by("-date", "-id"))
    expense_total = sum((Decimal(e.amount or 0) for e in expenses), Decimal("0"))

    transactions = list(
        FinancialTransaction.objects.filter(business=business, date__range=(date_from, date_to))
        .select_related("account")
        .order_by("-date", "-id")
    )
    money_in = sum((t.amount for t in transactions if t.transaction_type == FinancialTransaction.INCOME and not t.reversed), Decimal("0"))
    money_out = sum((t.amount for t in transactions if t.transaction_type == FinancialTransaction.OUTFLOW and not t.reversed), Decimal("0"))

    accounts = list(CashAccount.objects.filter(business=business, active=True).prefetch_related("transactions"))
    cash_balance = sum((account.balance for account in accounts), Decimal("0"))

    production_costs = list(
        ProductionCostSnapshot.objects.filter(business=business, production_date__range=(date_from, date_to))
        .select_related("finished_good", "production_batch")
        .prefetch_related("lines__raw_material")
        .order_by("-production_date", "-id")
    )
    batches = list(
        ProductionBatch.objects.filter(business=business, production_date__range=(date_from, date_to))
        .select_related("finished_good", "order")
        .order_by("-production_date", "-id")
    )
    quality_checks = list(
        ProductionQualityCheck.objects.filter(business=business, batch__production_date__range=(date_from, date_to))
        .select_related("batch__finished_good", "checked_by").order_by("-batch__production_date", "-id")
    )
    cost_snapshots = list(
        RawMaterialCostSnapshot.objects.filter(business=business, effective_date__range=(date_from, date_to))
        .select_related("raw_material", "purchase_order_item").order_by("-effective_date", "-id")
    )
    measurement_changes = list(
        RawMaterialMeasurementChange.objects.filter(
            business=business, created_at__date__range=(date_from, date_to)
        ).select_related("raw_material", "created_by").order_by("-created_at", "-id")
    )
    deliveries = list(
        DeliveryAssignment.objects.filter(business=business, created_at__date__range=(date_from, date_to))
        .select_related("intake", "driver", "quote", "origin", "provider_account")
        .prefetch_related("events")
    )
    logs_queryset = AuditLog.objects.filter(
        business=business,
        created_at__date__range=(date_from, date_to),
        model_name__in=_EXTERNAL_BUSINESS_ACTIVITY_MODELS,
    ).select_related("created_by")
    logs_total = logs_queryset.count()
    logs = list(logs_queryset if for_export else logs_queryset[:250])
    logs = [_prepare_audit_log(log) for log in logs]
    audit_queries = list(
        AuditQuery.objects.filter(business=business)
        .select_related("created_by", "assigned_to", "answered_by")
        .order_by("status", "-created_at", "-id")[:100]
    )

    return {
        "date_from": date_from,
        "date_to": date_to,
        "materials": materials,
        "products": products,
        "purchases": purchases,
        "sales": sales,
        "expenses": expenses,
        "transactions": transactions,
        "accounts": accounts,
        "production_costs": production_costs,
        "batches": batches,
        "quality_checks": quality_checks,
        "cost_snapshots": cost_snapshots,
        "measurement_changes": measurement_changes,
        "deliveries": deliveries,
        "logs": logs,
        "logs_total": logs_total,
        "logs_limited": not for_export and logs_total > len(logs),
        "audit_queries": audit_queries,
        "can_raise_audit_query": user_has_permission(request.user, business, "audit", "view"),
        "can_answer_audit_query": is_business_admin(request.user, business) or user_has_permission(request.user, business, "audit", "edit"),
        "audit_query_status_choices": AuditQuery.STATUS_CHOICES,
        "audit_query_severity_choices": AuditQuery.SEVERITY_CHOICES,
        "metrics": {
            "material_value": material_value,
            "product_value": product_value,
            "purchase_total": purchase_total,
            "supplier_payable": supplier_payable,
            "sales_total": sales_total,
            "receivables": receivables,
            "expense_total": expense_total,
            "money_in": money_in,
            "money_out": money_out,
            "net_cash_movement": money_in - money_out,
            "cash_balance": cash_balance,
        },
    }


@login_required
def audit_workspace(request):
    return render(request, "core/audit_workspace.html", _audit_dataset(request))


@login_required
def audit_query_create(request):
    if request.method != "POST" or not user_has_permission(request.user, request.business, "audit", "view"):
        return render(request, "403.html", status=403)
    subject = (request.POST.get("subject") or "").strip()[:160]
    message = (request.POST.get("message") or "").strip()
    if not subject or not message:
        messages.error(request, "Add a subject and query note before submitting.")
        return redirect("audit_workspace")
    severity = request.POST.get("severity") or AuditQuery.SEVERITY_MEDIUM
    valid_severities = {value for value, _ in AuditQuery.SEVERITY_CHOICES}
    if severity not in valid_severities:
        severity = AuditQuery.SEVERITY_MEDIUM
    query = AuditQuery.objects.create(
        business=request.business, created_by=request.user, module=(request.POST.get("module") or "General")[:40],
        record_label=(request.POST.get("record_label") or "")[:180], record_model=(request.POST.get("record_model") or "")[:120],
        record_id=(request.POST.get("record_id") or "")[:80], subject=subject, message=message[:4000],
        severity=severity,
    )
    audit(request.business, request.user, "audit_query_create", query, f"Audit query raised: {query.subject}", {"severity": query.severity, "module": query.module})
    messages.success(request, "Audit query submitted for management response.")
    return redirect("audit_workspace")


@login_required
def audit_query_update(request, pk):
    if request.method != "POST" or not (is_business_admin(request.user, request.business) or user_has_permission(request.user, request.business, "audit", "edit")):
        return render(request, "403.html", status=403)
    query = get_object_or_404(AuditQuery, pk=pk, business=request.business)
    status = request.POST.get("status") or query.status
    valid = {value for value, _ in AuditQuery.STATUS_CHOICES}
    if status not in valid:
        messages.error(request, "Choose a valid audit-query status.")
        return redirect("audit_workspace")
    query.status = status
    query.response = (request.POST.get("response") or query.response or "")[:4000]
    query.answered_by = request.user
    query.answered_at = timezone.now()
    if status == AuditQuery.STATUS_CLOSED:
        query.closed_at = timezone.now()
    query.save(update_fields=["status", "response", "answered_by", "answered_at", "closed_at", "updated_at"])
    audit(request.business, request.user, "audit_query_update", query, f"Audit query updated: {query.subject}", {"status": query.status})
    messages.success(request, "Audit query updated.")
    return redirect("audit_workspace")


@login_required
def audit_export_xlsx(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    data = _audit_dataset(request, for_export=True)
    business = request.business
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit summary"
    ws.append(["INPROFIC Audit Workspace", business.name])
    ws.append(["Period", f"{data['date_from']} to {data['date_to']}"])
    ws.append([])
    ws.append(["Metric", "Value"])
    for key, value in data["metrics"].items():
        ws.append([key.replace("_", " ").title(), float(value)])
    for cell in ws[4]:
        cell.font = Font(bold=True)

    sheets = [
        ("Materials", ["Name", "Category", "Purchase unit", "Usage unit", "Stock", "Unit cost", "Value"], [
            [m.name, m.get_category_display(), m.purchase_unit, m.usage_unit, float(m.stock or 0), float(m.cost_per_unit or 0), float(Decimal(m.stock or 0) * Decimal(m.cost_per_unit or 0))]
            for m in data["materials"]
        ]),
        ("Products", ["Name", "Category", "Source", "Unit", "Stock", "Sell price", "Estimated cost"], [
            [g.name, g.product_category.name if g.product_category else "", g.get_source_type_display(), g.unit, float(g.stock or 0), float(g.selling_price or 0), float(g.est_cost or 0)]
            for g in data["products"]
        ]),
        ("Product BOM", ["Product", "Line type", "Material", "Quantity per batch", "Usage unit", "Flexible"], [
            [g.name, "Recipe", row.raw_material.name, float(row.qty_per_batch), row.raw_material.usage_unit, bool(row.flexible_usage)]
            for g in data["products"] for row in g.recipe_items.all()
        ] + [
            [g.name, "Production input", row.raw_material.name, float(row.qty_per_batch), row.raw_material.usage_unit, False]
            for g in data["products"] for row in g.production_materials.all()
        ]),
        ("Procurement lines", ["Date", "PO", "Supplier", "Type", "Item", "Quantity", "Unit", "Unit cost", "Line total"], [
            [po.date.isoformat(), po.pk, po.supplier, item.item_type, item.item_name, float(item.qty), item.stock_unit, float(item.unit_cost), float(item.line_total)]
            for po in data["purchases"] for item in po.items.all()
        ]),
        ("Supplier payments", ["Date", "PO", "Supplier", "Method", "Reference", "Account", "Amount"], [
            [payment.date.isoformat(), po.pk, payment.supplier, payment.payment_method, payment.reference, payment.account.name if payment.account_id else "", float(payment.amount)]
            for po in data["purchases"] for payment in po.payments.all()
        ]),
        ("Procurement", ["Date", "PO", "Supplier", "Status", "Payment", "Total"], [
            [po.date.isoformat(), po.pk, po.supplier, po.status, po.payment_status, float(po.total)] for po in data["purchases"]
        ]),
        ("Sales", ["Date", "Sale", "Customer", "Source", "Payment state", "Total"], [
            [sale.date.isoformat(), sale.pk, sale.customer, sale.source, sale.transaction_type, float(sale.total)] for sale in data["sales"]
        ]),
        ("Sales lines", ["Date", "Sale", "Customer", "Product", "Units", "Unit price", "Discount", "Unit cost", "Line total"], [
            [sale.date.isoformat(), sale.pk, sale.customer, item.finished_good.name, float(item.total_units), float(item.price or 0), float(item.discount or 0), float(item.unit_cost or 0), float(item.line_total)]
            for sale in data["sales"] for item in sale.items.all()
        ]),
        ("Customer payments", ["Date", "Sale", "Customer", "Method", "Reference", "Account", "Amount"], [
            [payment.date.isoformat(), sale.pk, payment.customer, payment.payment_method, payment.reference, payment.account.name if payment.account_id else "", float(payment.amount)]
            for sale in data["sales"] for payment in sale.payments.all()
        ]),
        ("Cash accounts", ["Account", "Type", "Opening balance", "Current balance", "Active"], [
            [account.name, account.get_account_type_display(), float(account.opening_balance), float(account.balance), account.active]
            for account in data["accounts"]
        ]),
        ("Expenses", ["Date", "Category", "Description", "Vendor", "Payment", "Amount"], [
            [e.date.isoformat(), e.get_category_display(), e.description, e.vendor, e.payment_status, float(e.amount)] for e in data["expenses"]
        ]),
        ("Cash ledger", ["Date", "Type", "Category", "Description", "Account", "Amount", "Reversed"], [
            [t.date.isoformat(), t.transaction_type, t.category, t.description, t.account.name if t.account else "", float(t.amount), t.reversed] for t in data["transactions"]
        ]),
        ("Production costs", ["Date", "Product", "Batch", "Units", "Total cost", "Unit cost"], [
            [c.production_date.isoformat(), c.finished_good.name, c.batch_number, float(c.produced_units), float(c.total_cost), float(c.unit_cost)] for c in data["production_costs"]
        ]),
        ("Quality checks", ["Date", "Batch", "Product", "Result", "Checked by", "Checked at", "Notes", "Defects"], [
            [q.batch.production_date.isoformat(), q.batch.batch_number, q.batch.finished_good.name, q.status, str(q.checked_by or ""), q.checked_at.isoformat() if q.checked_at else "", q.notes, q.defects]
            for q in data["quality_checks"]
        ]),
        ("Measurement changes", ["Date", "Material", "Ratio", "Reason", "Old basis", "New basis"], [
            [c.created_at.isoformat(), c.raw_material.name, float(c.conversion_ratio), c.reason, str(c.old_measurement), str(c.new_measurement)] for c in data["measurement_changes"]
        ]),
        ("Deliveries", ["Created", "Order", "Customer", "Provider", "Driver", "Status", "Fee", "External ref"], [
            [d.created_at.isoformat(), d.intake.public_number, d.intake.customer_name, d.provider, d.driver.name if d.driver else "", d.status, float(d.intake.delivery_fee or 0), d.external_reference] for d in data["deliveries"]
        ]),
        ("Delivery events", ["When", "Order", "Delivery", "Status", "Note", "Metadata"], [
            [event.created_at.isoformat(), delivery.intake.public_number, str(delivery.public_id), event.status, redact_disabled_integrations(event.note), redact_disabled_integrations(str(event.metadata))]
            for delivery in data["deliveries"] for event in delivery.events.all()
        ]),
        ("Audit queries", ["When", "Status", "Severity", "Module", "Record", "Subject", "Query", "Response", "Raised by", "Answered by"], [
            [q.created_at.isoformat(), q.status, q.severity, q.module, q.record_label or q.record_id, q.subject, q.message, q.response, str(q.created_by or ""), str(q.answered_by or "")] for q in data["audit_queries"]
        ]),
        ("Audit trail", ["When", "Action", "Model", "Object", "Description", "Actor"], [
            [a.created_at.isoformat(), a.action, a.model_name, a.object_id, redact_disabled_integrations(a.description), str(a.created_by or "System")] for a in data["logs"]
        ]),
    ]
    for title, headers, rows in sheets:
        sheet = wb.create_sheet(title[:31])
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(max(12, max(len(str(c.value or "")) for c in column) + 2), 42)

    output = BytesIO()
    wb.save(output)
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{business.slug}-audit-{data["date_from"]}-{data["date_to"]}.xlsx"'
    return response
